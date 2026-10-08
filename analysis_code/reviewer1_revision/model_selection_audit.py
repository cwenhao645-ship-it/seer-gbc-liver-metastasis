"""Audit final-model designation using development-only out-of-fold predictions.

This script does not search hyperparameters, refit on temporal-test data, or write
row-level predictions. It reconstructs Tier 3 development OOF predictions using
the already selected hyperparameters and the original five-fold procedure, then
compares the documented development composite across candidate models.
"""

from __future__ import annotations

import argparse
import importlib.util
import json
import sys
from pathlib import Path
from types import ModuleType

import numpy as np
import openpyxl
import pandas as pd


def load_module(path: Path) -> ModuleType:
    spec = importlib.util.spec_from_file_location("locked_pipeline_audit", path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Cannot import locked pipeline: {path}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def load_tier3_parameters(path: Path) -> dict[str, dict]:
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    rows = list(worksheet.iter_rows(values_only=True))
    headers = list(rows[0])
    records = [dict(zip(headers, row)) for row in rows[1:] if row[0] == "Tier3"]
    return {
        str(record["Model"]): json.loads(str(record["Selected parameters"]))
        for record in records
    }


def build_args(module: ModuleType, raw: Path, output_root: Path) -> argparse.Namespace:
    parser = module.build_arg_parser()
    return parser.parse_args(
        [
            "--workdir",
            str(raw.parent),
            "--input",
            raw.name,
            "--outdir",
            str(output_root / "unused_pipeline_output"),
            "--dev-years-start",
            "2010",
            "--dev-years-end",
            "2018",
            "--test-years-start",
            "2019",
            "--test-years-end",
            "2022",
            "--cv-folds",
            "5",
            "--random-state",
            "42",
            "--n-jobs-search",
            "1",
            "--n-jobs-model",
            "1",
            "--selection-strategy",
            "balanced",
            "--tree-objective",
            "balanced",
            "--tree-calibration",
            "both",
            "--threshold-rule",
            "sensitivity80",
            "--target-sensitivity",
            "0.8",
            "--skip-shap",
        ]
    )


def fit_locked_candidate(
    module: ModuleType,
    model_name: str,
    selected_params: dict,
    x_dev: pd.DataFrame,
    y_dev: np.ndarray,
    feature_columns: list[str],
    numeric_columns: list[str],
    categorical_columns: list[str],
    args: argparse.Namespace,
) -> tuple[np.ndarray, str]:
    positives = max(1, int(y_dev.sum()))
    negatives = max(1, int(len(y_dev) - y_dev.sum()))
    scale_pos_weight = negatives / positives

    if model_name == "CatBoost":
        parameters = {
            key: value
            for key, value in selected_params.items()
            if key not in {"probability_calibration", "tree_balanced_oof_score"}
        }
        estimator = module.make_catboost_estimator(
            parameters, scale_pos_weight, args, seed_offset=9999
        )
        frame, medians, cat_indices = module.prepare_catboost_frame(
            x_dev, feature_columns, numeric_columns, categorical_columns
        )
        estimator.fit(frame, y_dev, cat_features=cat_indices)
        raw_oof = module.oof_predict_model(
            estimator,
            x_dev,
            y_dev,
            model_name,
            feature_columns,
            numeric_columns,
            categorical_columns,
            args,
            medians,
            cat_indices,
        )
        calibrated, _, method, _, _ = module.apply_development_only_calibration(
            y_dev, raw_oof, raw_oof, args
        )
        return calibrated, method

    estimator, _ = module.make_sklearn_estimator(
        model_name,
        numeric_columns,
        categorical_columns,
        args,
        scale_pos_weight,
    )
    if model_name in {"Linear Logistic", "Spline Logistic"}:
        estimator.set_params(clf__C=float(selected_params["C"]))
    elif model_name == "XGBoost":
        model_params = {
            f"clf__{key}": value
            for key, value in selected_params.items()
            if key not in {"probability_calibration", "tree_balanced_oof_score"}
        }
        estimator.set_params(**model_params)
    else:
        raise ValueError(model_name)

    estimator.fit(x_dev, y_dev)
    raw_oof = module.oof_predict_model(
        estimator,
        x_dev,
        y_dev,
        model_name,
        feature_columns,
        numeric_columns,
        categorical_columns,
        args,
    )
    if model_name == "XGBoost":
        calibrated, _, method, _, _ = module.apply_development_only_calibration(
            y_dev, raw_oof, raw_oof, args
        )
        return calibrated, method
    return raw_oof, "none"


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--pipeline", type=Path, required=True)
    parser.add_argument("--raw", type=Path, required=True)
    parser.add_argument("--table-s4", type=Path, required=True)
    parser.add_argument("--outdir", type=Path, required=True)
    parsed = parser.parse_args()

    parsed.outdir.mkdir(parents=True, exist_ok=True)
    module = load_module(parsed.pipeline)
    args = build_args(module, parsed.raw, parsed.outdir)
    raw = pd.read_csv(parsed.raw, low_memory=False)
    main_cohort, flow = module.build_main_cohort(raw, args)
    tier3 = module.make_tier_datasets(main_cohort)["Tier3"]
    development = tier3[
        tier3["Temporal_Split"].astype(str).str.startswith("Development")
    ].copy()
    feature_columns = module.get_feature_columns("Tier3", tier3)
    numeric_columns, categorical_columns = module.split_num_cat(
        development, feature_columns
    )
    x_dev = development[feature_columns].copy()
    y_dev = development[module.TARGET_COL].astype(int).to_numpy()
    parameters = load_tier3_parameters(parsed.table_s4)

    rows = []
    for model_name in [
        "Linear Logistic",
        "Spline Logistic",
        "XGBoost",
        "CatBoost",
    ]:
        probabilities, calibration = fit_locked_candidate(
            module,
            model_name,
            parameters[model_name],
            x_dev,
            y_dev,
            feature_columns,
            numeric_columns,
            categorical_columns,
            args,
        )
        threshold = module.select_threshold(
            y_dev,
            probabilities,
            rule="sensitivity80",
            target_sens=0.8,
        )
        metrics = module.compute_metrics(y_dev, probabilities, threshold)
        documented_score = module.tree_multi_metric_score(y_dev, probabilities, args)
        no_slope_score = (
            2.00 * float(metrics["PR-AUC"])
            + 0.75 * float(metrics["ROC-AUC"])
            + 0.35 * float(metrics["F1"])
            - 1.25 * float(metrics["Brier"])
        )
        rows.append(
            {
                "Model": model_name,
                "Development_N": len(y_dev),
                "Development_events": int(y_dev.sum()),
                "Calibration": calibration,
                "OOF_ROC_AUC": metrics["ROC-AUC"],
                "OOF_PR_AUC": metrics["PR-AUC"],
                "OOF_Brier": metrics["Brier"],
                "OOF_F1": metrics["F1"],
                "OOF_calibration_intercept": metrics["Calibration intercept"],
                "OOF_calibration_slope": metrics["Calibration slope"],
                "OOF_threshold": threshold,
                "Documented_composite_with_slope_penalty": documented_score,
                "Composite_without_slope_penalty": no_slope_score,
            }
        )

    result = pd.DataFrame(rows).sort_values(
        "Documented_composite_with_slope_penalty", ascending=False
    )
    csv_path = parsed.outdir / "development_oof_model_selection_audit.csv"
    result.to_csv(csv_path, index=False)
    winner = str(result.iloc[0]["Model"])
    catboost = result[result["Model"].eq("CatBoost")].iloc[0]
    report = {
        "audit_scope": "Tier 3 fixed-hyperparameter development-only OOF adjudication",
        "raw_records": int(flow["raw_n"]),
        "development_n": int(len(y_dev)),
        "development_events": int(y_dev.sum()),
        "documented_composite_winner": winner,
        "catboost_locked_threshold_reconstructed": float(catboost["OOF_threshold"]),
        "expected_locked_threshold": 0.08367640386406207,
        "threshold_absolute_difference": abs(
            float(catboost["OOF_threshold"]) - 0.08367640386406207
        ),
        "row_level_predictions_written": False,
        "output_csv": str(csv_path),
    }
    json_path = parsed.outdir / "model_selection_audit.json"
    json_path.write_text(
        json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8"
    )
    print(json.dumps(report, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
